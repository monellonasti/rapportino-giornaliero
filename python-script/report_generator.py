"""Rapportino vendite in Excel, con lo stesso layout dell'app Shopify."""
import re
import secrets
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Optional
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.pagebreak import Break

# Stessa palette dell'app (app/export/theme.ts).
BRAND = '173B36'
ON_BRAND = 'C9DDD6'
ACCENT = '174F43'
ACCENT_SOFT = 'E6F0EC'
INK = '16241F'
INK2 = '3F504A'
MUTED = '5F6F6A'
LINE = 'E2E8E5'
LINE_STRONG = 'C9D3CF'
SOFT = 'F1F5F3'
WHITE = 'FFFFFF'
WARN = '7A4D00'
WARN_LINE = 'F0D08A'

FONT = 'Aptos'
MONEY = '#,##0.00 "€";-#,##0.00 "€"'
MONEY_OR_DASH = '#,##0.00 "€";-#,##0.00 "€";"–"'
PERCENT = '0.0%'
PERCENT_OR_DASH = '0.0%;-0.0%;"–"'
LAST = 8  # colonne A:H
WIDTHS = [42, 7, 16, 14, 14, 12, 11, 15]
HEADS = ['Prodotto', 'Q.tà', 'Prezzo confronto', 'Prezzo vendita', 'Tot. confronto', 'Sconto €', 'Sconto %', 'Totale vendita']
LINE_FORMATS = [None, '0', MONEY, MONEY, MONEY, MONEY_OR_DASH, PERCENT_OR_DASH, MONEY]
CARDS = [(1, 1), (2, 5), (6, 8)]
# Altezza utile, in punti, di una pagina A4 orizzontale con i margini del foglio.
PAGE_HEIGHT = 490
GIORNI = ['lunedì', 'martedì', 'mercoledì', 'giovedì', 'venerdì', 'sabato', 'domenica']
MESI = ['gennaio', 'febbraio', 'marzo', 'aprile', 'maggio', 'giugno', 'luglio', 'agosto',
        'settembre', 'ottobre', 'novembre', 'dicembre']


def money(value):
    if value is None or value == '':
        return Decimal('0.00')
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def get_amount(money_set):
    try:
        return money(money_set['shopMoney']['amount'])
    except Exception:
        return Decimal('0.00')


def euro(value):
    return f'{money(value):,.2f} €'.replace(',', 'X').replace('.', ',').replace('X', '.')


def percent(value):
    return f'{value * 100:.1f}%'.replace('.', ',')


def plural(n, one, many):
    return f'{n} {one if n == 1 else many}'


def normalize_channel(source_name):
    if not source_name:
        return 'Online'
    src = source_name.lower()
    if 'pos' in src:
        return 'POS'
    if 'web' in src or 'online' in src or src == 'shopify_draft_order':
        return 'Online'
    return str(source_name)


def payment_label(gateway):
    """Etichetta leggibile del gateway, come nell'app."""
    g = str(gateway).strip().lower()
    if not g:
        return 'Non specificato'
    if re.search(r'delivery|\bcod\b|contrassegno', g):
        return 'Contrassegno'
    if 'cash' in g or 'contant' in g:
        return 'Contanti'
    if 'gift' in g:
        return 'Gift card'
    if 'paypal' in g:
        return 'PayPal'
    if re.search(r'bank|bonifico|transfer', g):
        return 'Bonifico'
    if re.search(r'shopify.?payments|stripe|card|carta|visa|mastercard|maestro|amex|\bpos\b', g):
        return 'Carta'
    if g == 'manual':
        return 'Manuale'
    return str(gateway).strip()


def report_title(report_date):
    day = datetime.strptime(report_date, '%Y-%m-%d')
    # "del venerdì", ma "della domenica" (femminile).
    article = 'della' if day.weekday() == 6 else 'del'
    return f'Rapportino {article} {GIORNI[day.weekday()]} {day.day} {MESI[day.month - 1]} {day.year}'


@dataclass
class Line:
    name: str
    qty: int
    compare: Decimal
    sold: Decimal
    total: Decimal
    discount: Decimal
    discount_pct: float


@dataclass
class Sale:
    name: str
    time: str
    channel: str
    payment: str
    total: Decimal
    items: int
    compare_total: Decimal
    # Merce dopo tutti gli sconti; sconto e percentuale includono lo sconto sull'ordine.
    merchandise: Decimal
    discount: Decimal
    discount_pct: float
    order_discount: Decimal
    shipping: Decimal
    # Resto del totale ordine oltre merce e spedizione (tasse escluse, mance…).
    other: Decimal
    lines: list


@dataclass
class Report:
    date: str
    title: str
    shop_name: str
    timezone: str
    status: str  # 'bozza' | 'definitivo'
    generated_at: datetime
    closed_by: Optional[str]
    sales: list
    kpis: list
    summary: list
    payments: list
    payments_total: Decimal

    @property
    def final(self):
        return self.status == 'definitivo'

    def status_label(self):
        if self.final:
            by = f' da {self.closed_by}' if self.closed_by else ''
            return f'Definitivo · chiuso il {self.generated_at:%d/%m/%Y} alle {self.generated_at:%H:%M}{by}'
        return 'Bozza · non valida come chiusura'


def build_sale(order, zone):
    lines = []
    for edge in order.get('lineItems', {}).get('edges', []):
        item = edge['node']
        qty = int(item.get('quantity') or 0)
        # Prezzo venduto al netto degli sconti di riga; il prezzo di listino è il "prezzo di confronto".
        sold = get_amount(item.get('discountedUnitPriceSet')) or get_amount(item.get('originalUnitPriceSet'))
        variant = item.get('variant') or {}
        compare = money(variant.get('compareAtPrice')) if variant.get('compareAtPrice') else Decimal('0')
        if compare <= 0:
            compare = sold
        total = sold * qty
        listed = compare * qty
        lines.append(Line(item.get('name', ''), qty, compare, sold, total, listed - total,
                          float((listed - total) / listed) if listed else 0.0))
    lines_total = sum((l.total for l in lines), Decimal('0'))
    compare_total = sum((l.compare * l.qty for l in lines), Decimal('0'))
    # I prezzi di riga escludono gli sconti sull'ordine, il subtotale Shopify li comprende.
    merchandise = get_amount(order['subtotalPriceSet']) if order.get('subtotalPriceSet') else lines_total
    total = get_amount(order.get('totalPriceSet'))
    shipping = get_amount(order.get('totalShippingPriceSet'))
    labels = []
    for gateway in order.get('paymentGatewayNames') or []:
        label = payment_label(gateway)
        if label not in labels:
            labels.append(label)
    created = order.get('createdAt')
    time = datetime.fromisoformat(created.replace('Z', '+00:00')).astimezone(zone).strftime('%H:%M') if created else ''
    return Sale(
        name=order.get('name', ''),
        time=time,
        channel=normalize_channel(order.get('sourceName')),
        payment=' + '.join(labels) or 'Non specificato',
        total=total,
        items=sum(l.qty for l in lines),
        compare_total=compare_total,
        merchandise=merchandise,
        discount=compare_total - merchandise,
        discount_pct=float((compare_total - merchandise) / compare_total) if compare_total else 0.0,
        order_discount=lines_total - merchandise,
        shipping=shipping,
        other=total - merchandise - shipping,
        lines=lines,
    )


def build_report(orders, report_date, shop_name, timezone, status='bozza', closed_by=None, now=None):
    zone = ZoneInfo(timezone)
    valid = [o for o in orders if not o.get('test') and not o.get('cancelledAt')]
    sales = [build_sale(o, zone) for o in valid]
    total = sum((s.total for s in sales), Decimal('0'))
    tax = sum((get_amount(o.get('totalTaxSet')) for o in valid), Decimal('0'))
    discounts = sum((get_amount(o.get('totalDiscountsSet')) for o in valid), Decimal('0'))
    merchandise = sum((s.merchandise for s in sales), Decimal('0'))
    compare_total = sum((s.compare_total for s in sales), Decimal('0'))
    shipping = sum((s.shipping for s in sales), Decimal('0'))
    items = sum(s.items for s in sales)
    count = len(sales)
    compare_discount = compare_total - merchandise
    rate = float(compare_discount / compare_total) if compare_total else 0.0
    average = money(total / count) if count else Decimal('0.00')
    share = (lambda channel: sum(1 for s in sales if s.channel == channel) / count if count else 0.0)
    cash = sum((s.total for s in sales if s.payment == 'Contanti'), Decimal('0'))
    mixed = sum(1 for s in sales if 'Contanti' in s.payment.split(' + ') and s.payment != 'Contanti')
    payments = {}
    for s in sales:
        payments[s.payment] = payments.get(s.payment, Decimal('0')) + s.total
    return Report(
        date=report_date,
        title=report_title(report_date),
        shop_name=shop_name,
        timezone=timezone,
        status=status,
        generated_at=now or datetime.now(zone),
        closed_by=closed_by,
        sales=sales,
        kpis=[
            ('Vendite IVA inclusa', total, 'money', 'Spedizioni incluse', False),
            ('Vendite IVA esclusa', total - tax, 'money', f'IVA {euro(tax)}', False),
            ('Sconto totale', compare_discount, 'money', f'{percent(rate)} sul prezzo confronto', False),
            ('Numero vendite', count, 'number', plural(items, 'articolo venduto', 'articoli venduti'), False),
            ('Scontrino medio', average, 'money', 'IVA inclusa', False),
            ('Incasso contanti', cash, 'money',
             f'Esclusi {plural(mixed, "pagamento misto", "pagamenti misti")}' if mixed else 'Vendite pagate in contanti',
             True),
        ],
        # Voci di vendita nell'ordine del rapportino originale.
        summary=[
            ('Totale vendite IVA inclusa', total, 'money', True),
            ('Totale vendite IVA esclusa', total - tax, 'money', True),
            ('Sconto confronto totale', compare_discount, 'money', True),
            ('Sconto confronto medio', rate, 'percent', False),
            ('Sconti applicati Shopify', discounts, 'money', False),
            ('Totale merce', merchandise, 'money', False),
            ('Totale spedizione', shipping, 'money', False),
            ('Numero vendite', count, 'number', False),
            ('Percentuale Online', share('Online'), 'percent', False),
            ('Percentuale POS', share('POS'), 'percent', False),
            ('Articoli venduti', items, 'number', False),
            ('Scontrino medio', average, 'money', False),
        ],
        payments=sorted(payments.items(), key=lambda item: item[1], reverse=True),
        payments_total=total,
    )


# --- Scrittura Excel -------------------------------------------------------------

def _font(**kw):
    return Font(name=FONT, size=kw.pop('size', 10), color=kw.pop('color', INK), **kw)


def _side(color, style='thin'):
    return Side(style=style, color=color)


def _num_format(kind):
    return MONEY if kind == 'money' else PERCENT if kind == 'percent' else '0'


def _value(value, kind):
    return float(value) if kind == 'money' else value


def _paint(ws, row, first, last, font=None, fill=None, align=None, num_format=None, border=None):
    for col in range(first, last + 1):
        cell = ws.cell(row, col)
        if font is not None:
            cell.font = font
        if fill is not None:
            cell.fill = PatternFill('solid', fgColor=fill)
        if align is not None:
            cell.alignment = align
        if num_format is not None:
            cell.number_format = num_format
        if border:
            old = cell.border
            cell.border = Border(left=border.get('left', old.left), right=border.get('right', old.right),
                                 top=border.get('top', old.top), bottom=border.get('bottom', old.bottom))


def _put(ws, row, col, value, last=None, font=None, **look):
    last = last or col
    if last > col:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=last)
    ws.cell(row, col).value = value
    _paint(ws, row, col, last, font=font or _font(), **look)


def _height(ws, row, value):
    ws.row_dimensions[row].height = value


def _wrap_height(text, chars_per_line, line_height=13.5):
    return 6 + line_height * max(1, -(-len(text) // chars_per_line))


def _align(col, **extra):
    base = dict(horizontal='left' if col == 1 else 'center' if col == 2 else 'right',
                indent=1 if col == 1 else 0, vertical='center')
    base.update(extra)
    return Alignment(**base)


def _rich(*parts):
    """Testo con più stili nella stessa cella: (testo, dimensione, grassetto, colore)."""
    return CellRichText(*[TextBlock(InlineFont(rFont=FONT, sz=size, b=bold, color=color), text)
                          for text, size, bold, color in parts])


def _section(ws, row, title, aside=''):
    rule = {'bottom': _side(BRAND, 'medium')}
    _put(ws, row, 1, title, font=_font(size=13, bold=True), align=Alignment(vertical='bottom'), border=rule)
    _put(ws, row, 2, aside, LAST, font=_font(size=9, color=MUTED),
         align=Alignment(horizontal='right', vertical='bottom'), border=rule)
    _height(ws, row, 24)
    _height(ws, row + 1, 8)
    return row + 2


def _header(ws, report):
    _put(ws, 1, 1, report.title, LAST - 2, font=_font(size=18, bold=True, color=WHITE), fill=BRAND,
         align=Alignment(vertical='center', indent=1))
    # Stato in alto a destra: bozza in ambra, definitivo in chiaro sul verde.
    _put(ws, 1, LAST - 1, 'DEFINITIVO' if report.final else 'BOZZA', LAST,
         font=_font(size=12, bold=True, color=ON_BRAND if report.final else WARN_LINE), fill=BRAND,
         align=Alignment(horizontal='right', vertical='center', indent=1))
    _height(ws, 1, 36)
    meta = (f'   ·   Fuso orario {report.timezone}   ·   Generato il {report.generated_at:%d/%m/%Y} '
            f'alle {report.generated_at:%H:%M}   ·   {report.status_label()}')
    _put(ws, 2, 1, _rich((report.shop_name, 10.5, True, WHITE), (meta, 9, False, ON_BRAND)), LAST,
         fill=BRAND, align=Alignment(vertical='top', indent=1))
    _height(ws, 2, 22)
    _height(ws, 3, 12)
    return 4


def _kpis(ws, report, top):
    for i, (label, value, kind, hint, highlight) in enumerate(report.kpis):
        first, last = CARDS[i % len(CARDS)]
        row = top + (i // len(CARDS)) * 4
        fill = ACCENT_SOFT if highlight else SOFT
        _put(ws, row, first, label, last, font=_font(size=9, bold=True, color=MUTED), fill=fill,
             align=Alignment(vertical='bottom', indent=1))
        _put(ws, row + 1, first, _value(value, kind), last,
             font=_font(size=18, bold=True, color=ACCENT if highlight else INK), fill=fill,
             num_format=_num_format(kind), align=Alignment(horizontal='left', vertical='center', indent=1))
        _put(ws, row + 2, first, hint, last, font=_font(size=9, color=MUTED), fill=fill,
             align=Alignment(vertical='top', indent=1))
        # Bordi bianchi spessi tra i riquadri per staccarli come schede.
        for r in range(row, row + 3):
            if first > 1:
                _paint(ws, r, first, first, border={'left': _side(WHITE, 'thick')})
            if last < LAST:
                _paint(ws, r, last, last, border={'right': _side(WHITE, 'thick')})
        for offset, value in enumerate((19, 28, 19, 8)):
            _height(ws, row + offset, value)
    return top + -(-len(report.kpis) // len(CARDS)) * 4 + 1


def _sale(ws, sale, top):
    r = top
    band = {'fill': ACCENT_SOFT}
    _put(ws, r, 1, _rich((sale.name, 11, True, INK), (f'    {sale.time}   ·   {sale.channel}   ·   {sale.payment}', 10, False, INK2)),
         align=Alignment(vertical='center', indent=1), **band)
    _put(ws, r, 2, 'Totale ordine', LAST - 1, font=_font(size=9, color=MUTED),
         align=Alignment(horizontal='right', vertical='center'), **band)
    _put(ws, r, LAST, float(sale.total), font=_font(size=11, bold=True), num_format=MONEY,
         align=Alignment(vertical='center'), **band)
    _height(ws, r, 24)
    r += 1
    for i, head in enumerate(HEADS):
        _put(ws, r, i + 1, head, font=_font(size=8.5, bold=True, color=MUTED),
             align=_align(i + 1, wrap_text=True, vertical='bottom'), border={'bottom': _side(LINE_STRONG)})
    _height(ws, r, 24)
    r += 1
    for line in sale.lines:
        values = [line.name, line.qty, float(line.compare), float(line.sold), float(line.compare * line.qty),
                  float(line.discount), line.discount_pct, float(line.total)]
        for i, value in enumerate(values):
            _put(ws, r, i + 1, value, font=_font(bold=i == LAST - 1), num_format=LINE_FORMATS[i],
                 align=_align(i + 1, vertical='top', wrap_text=i == 0), border={'bottom': _side(LINE)})
        _height(ws, r, _wrap_height(line.name, 46))
        r += 1
    if len(sale.lines) > 1 or sale.order_discount:
        label = [('Totale merce', 10, True, INK)]
        if sale.order_discount:
            label.append((f'   incl. sconto ordine {euro(sale.order_discount)}', 8.5, False, MUTED))
        values = [_rich(*label), sale.items, None, None, float(sale.compare_total), float(sale.discount),
                  sale.discount_pct, float(sale.merchandise)]
        for i, value in enumerate(values):
            _put(ws, r, i + 1, value, font=_font(bold=True), fill=SOFT, num_format=LINE_FORMATS[i],
                 align=_align(i + 1), border={'top': _side(LINE_STRONG)})
        _height(ws, r, 20)
        r += 1
    # Voci che riconciliano la merce con il totale dell'ordine.
    extra = [('Spedizione', sale.shipping)] if sale.shipping else []
    if sale.other:
        extra.append(("Altre voci dell'ordine", sale.other))
    for label, value in extra:
        rule = {'bottom': _side(LINE)}
        _put(ws, r, 1, label, LAST - 1, font=_font(color=INK2), align=Alignment(indent=1, vertical='center'), border=rule)
        _put(ws, r, LAST, float(value), font=_font(bold=True), num_format=MONEY, align=Alignment(vertical='center'),
             border=rule)
        _height(ws, r, 18)
        r += 1
    _height(ws, r, 12)
    return r + 1


def _panel_title(ws, row, first, last, title):
    _put(ws, row, first, title, last, font=_font(size=11, bold=True), fill=SOFT,
         align=Alignment(vertical='center', indent=1))


def _recap(ws, report, top):
    rule = {'bottom': _side(LINE)}
    # Colonna sinistra: voci di vendita. Colonna destra: incasso per metodo di pagamento.
    _panel_title(ws, top, 1, 3, 'Vendite')
    for i, (label, value, kind, strong) in enumerate(report.summary):
        row = top + 1 + i
        font = _font(bold=strong, color=INK if strong else INK2)
        _put(ws, row, 1, label, font=font, align=Alignment(vertical='center', indent=1), border=rule)
        _put(ws, row, 2, _value(value, kind), 3, font=font, num_format=_num_format(kind),
             align=Alignment(horizontal='right', vertical='center'), border=rule)
    r = top
    _panel_title(ws, r, 5, 8, 'Incasso per metodo di pagamento')
    r += 1
    heads = {'font': _font(size=8.5, bold=True, color=MUTED), 'border': {'bottom': _side(LINE_STRONG)}}
    _put(ws, r, 5, 'Metodo', 6, align=Alignment(indent=1, vertical='center'), **heads)
    _put(ws, r, 7, 'Importo', align=Alignment(horizontal='right', vertical='center'), **heads)
    _put(ws, r, 8, '% incasso', align=Alignment(horizontal='right', vertical='center'), **heads)
    r += 1
    if not report.payments:
        _put(ws, r, 5, 'Nessuna vendita nel giorno', 8, font=_font(italic=True, color=MUTED),
             align=Alignment(indent=1), border=rule)
        r += 1
    for label, amount in report.payments:
        share = float(amount / report.payments_total) if report.payments_total else 0.0
        _put(ws, r, 5, label, 6, align=Alignment(indent=1, vertical='center'), border=rule)
        _put(ws, r, 7, float(amount), num_format=MONEY, align=Alignment(vertical='center'), border=rule)
        _put(ws, r, 8, share, num_format=PERCENT, align=Alignment(vertical='center'), border=rule)
        r += 1
    totals = {'font': _font(bold=True), 'border': {'top': _side(LINE_STRONG)}}
    _put(ws, r, 5, 'Totale', 6, align=Alignment(indent=1, vertical='center'), **totals)
    _put(ws, r, 7, float(report.payments_total), num_format=MONEY, align=Alignment(vertical='center'), **totals)
    _put(ws, r, 8, 1 if report.payments_total else 0, num_format=PERCENT, align=Alignment(vertical='center'), **totals)
    r += 1
    end = max(top + len(report.summary) + 1, r)
    for row in range(top, end):
        _height(ws, row, 19)
    _height(ws, end, 12)
    return end + 1


def _page_breaks(ws, starts, end):
    """Interruzioni manuali prima dei blocchi che non entrano nella pagina corrente."""
    used = 0
    for i, start in enumerate(starts):
        stop = starts[i + 1] if i + 1 < len(starts) else end
        block = sum(ws.row_dimensions[r].height or 15 for r in range(start, stop))
        if used > 0 and used + block > PAGE_HEIGHT:
            ws.row_breaks.append(Break(id=start - 1))
            used = 0
        used = (used + block) % PAGE_HEIGHT


def generate_excel(report, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Rapportino'
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    # Margine alto più ampio: lascia spazio alla scritta "BOZZA" nell'intestazione di pagina.
    ws.page_margins.top = 0.7
    ws.page_margins.bottom = 0.6
    ws.page_margins.header = 0.25
    ws.page_margins.footer = 0.3
    for i, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    # "&" è un codice di formato nell'intestazione Excel: va raddoppiato.
    footer = (f'Rapportino del {datetime.strptime(report.date, "%Y-%m-%d"):%d/%m/%Y} · {report.shop_name} · '
              f'{report.status_label()}').replace('&', '&&')
    ws.oddFooter.left.text = footer
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = 'Pagina &P di &N'
    ws.oddFooter.right.size = 8
    if not report.final:
        # Sulla carta la bozza porta "BOZZA" in alto su ogni pagina.
        ws.oddHeader.center.text = 'BOZZA'
        ws.oddHeader.center.size = 18
        ws.oddHeader.center.font = f'{FONT},Bold'
        ws.oddHeader.center.color = WARN

    r = _kpis(ws, report, _header(ws, report))
    blocks = [1, r]  # prime righe dei blocchi da non spezzare in stampa
    items = sum(s.items for s in report.sales)
    r = _section(ws, r, 'Dettaglio vendite',
                 f'{plural(len(report.sales), "vendita", "vendite")}  ·  {plural(items, "articolo", "articoli")}')
    if not report.sales:
        _put(ws, r, 1, 'Nessuna vendita nella data scelta.', LAST, font=_font(italic=True, color=MUTED),
             align=Alignment(horizontal='center', vertical='center'))
        _height(ws, r, 28)
        r += 2
    for i, sale in enumerate(report.sales):
        if i:
            blocks.append(r)
        r = _sale(ws, sale, r)
    blocks.append(r)
    r = _recap(ws, report, _section(ws, r, 'Riepilogo giornaliero'))
    _put(ws, r, 1, 'Generato automaticamente da Shopify Rapportino · verifica sempre resi, cambi merce e pagamenti '
                   'misti prima della chiusura definitiva.', LAST, font=_font(size=8.5, italic=True, color=MUTED),
         align=Alignment(horizontal='center', vertical='center'))
    _height(ws, r, 18)
    ws.print_area = f'A1:H{r}'
    _page_breaks(ws, blocks, r + 1)

    if report.final:
        # Definitivo: foglio e struttura bloccati con password casuale non salvata.
        password = secrets.token_urlsafe(24)
        ws.protection.sheet = True
        ws.protection.password = password
        wb.security = WorkbookProtection(lockStructure=True)
        wb.security.set_workbook_password(password)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
